"""Build human-verified gold labels for `cli eval`, starting from the extractor's own output.

    python tools/gold.py draft dataset/ --n 50          # sample, extract, write gold_draft/review.xlsx
    (a person checks each document in review.xlsx and marks it Verified = Y)
    python tools/gold.py build                          # verified docs -> gold/<name>.json
    python -m vaxextract.cli eval --pred gold_draft/preds --gold gold

`draft` keeps the --full predictions in gold_draft/preds, so the eval above scores exactly the run the
reviewer corrected. Reads .env like the CLI.
"""
from __future__ import annotations

import argparse
import json
import random
import sys
import time
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from vaxextract.pipeline import Config, extract_file, load_dotenv

EXTS = {".pdf", ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
DEFAULT_TARGETS = "Bordetella;Canine Influenza;DHLPP / FVRCP;Parvo;Rabies"
DOC_COLS = ["Doc", "File", "Type", "Pets (as printed, ';'-separated)", "Verified (Y)", "Notes", "Extractor review reasons"]
EV_COLS = ["Doc", "Pet", "Requested name", "Printed name", "Date given (yyyy-mm-dd)", "Due date (yyyy-mm-dd)",
           "Status", "Evidence (OCR line)"]
STATUSES = ["OK", "Fixed", "Delete", "Added"]


def sample(folder: Path, n: int, seed: int) -> list[Path]:
    """`n` files spread over the file types in proportion to how often each occurs (jpg/jpeg count as one)."""
    by_type = defaultdict(list)
    for f in sorted(folder.iterdir()):
        if f.suffix.lower() in EXTS:
            by_type[f.suffix.lower().replace(".jpeg", ".jpg")].append(f)
    total = sum(map(len, by_type.values()))
    rng = random.Random(seed)
    quota = {t: max(1, round(n * len(fs) / total)) for t, fs in by_type.items()}
    while sum(quota.values()) > n:
        quota[max(quota, key=quota.get)] -= 1
    while sum(quota.values()) < n:
        quota[max(by_type, key=lambda t: len(by_type[t]) - quota[t])] += 1
    picked = [f for t, fs in by_type.items() for f in rng.sample(fs, min(quota[t], len(fs)))]
    return sorted(picked)


def draft(a) -> int:
    load_dotenv(Path.cwd() / ".env")
    cfg = Config.from_env()
    targets = [t.strip() for t in a.targets.split(";") if t.strip()]
    out, preds = Path(a.out), Path(a.out) / "preds"
    preds.mkdir(parents=True, exist_ok=True)
    files = sample(Path(a.folder), a.n, a.seed)

    wb = openpyxl.Workbook()
    docs = wb.active
    docs.title = "Docs"
    evs = wb.create_sheet("Events")
    bold = Font(bold=True)
    for ws, cols in ((docs, DOC_COLS), (evs, EV_COLS)):
        ws.append(cols)
        for c in ws[1]:
            c.font = bold
        ws.freeze_panes = "A2"
    flag = PatternFill("solid", fgColor="FBE3C8")

    for i, f in enumerate(files, 1):
        pp = preds / (f.stem + ".json")
        t0 = time.time()
        if pp.exists():  # resume an interrupted run without paying for the call again
            res = json.loads(pp.read_text(encoding="utf-8"))
        else:
            try:
                res = json.loads(extract_file(f, cfg, targets=targets).model_dump_json())
            except Exception as e:  # keep the batch going; the doc is still listed for manual labelling
                print(f"[{i}/{len(files)}] {f.name}: ERROR {e}", file=sys.stderr)
                res = {"data": {"pets": []}, "meta": {"reviewReasons": [f"extraction failed: {e}"]}}
            pp.write_text(json.dumps(res, indent=2), encoding="utf-8")
        pets = res["data"]["pets"]
        reasons = "; ".join(res.get("meta", {}).get("reviewReasons", []))
        docs.append([f.stem, f.name, f.suffix.lower().lstrip("."), "; ".join(p.get("name") or "" for p in pets),
                     None, None, reasons])
        link = docs.cell(docs.max_row, 2)
        link.hyperlink, link.style = f.resolve().as_uri(), "Hyperlink"
        if reasons:
            docs.cell(docs.max_row, 7).fill = flag
        for p in pets:
            for e in p.get("events", []):
                if e.get("target"):
                    evs.append([f.stem, p.get("name"), e["target"], e.get("rawName"), e.get("dateAdministered"),
                                e.get("dueDate"), None, e.get("evidence")])
                    if e.get("needsHumanReview"):
                        for c in evs[evs.max_row]:
                            c.fill = flag
        print(f"[{i}/{len(files)}] {f.name}: {sum(len(p.get('events', [])) for p in pets)} events "
              f"({time.time() - t0:.1f} s)")

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    evs.add_data_validation(dv)
    dv.add("G2:G5000")
    yn = DataValidation(type="list", formula1='"Y"', allow_blank=True)
    docs.add_data_validation(yn)
    yn.add("E2:E500")
    for ws, widths in ((docs, [38, 44, 6, 30, 12, 30, 50]), (evs, [38, 18, 18, 40, 20, 20, 10, 60])):
        for col, w in zip("ABCDEFGH", widths):
            ws.column_dimensions[col].width = w
    info = wb.create_sheet("How to review")
    for line in [
        "For each document in Docs: open the file (link), then check its rows in Events against the document.",
        f"Requested names: {'; '.join(targets)}. Only vaccines that cover one of these belong in Events.",
        "Events: a correct row -> Status OK. Wrong date or name -> correct the cell, Status Fixed.",
        "A row that should not exist -> Status Delete. A missed vaccine -> add a row (Doc, Pet, Requested name, dates), Status Added.",
        "One row per (printed vaccine line x requested name it covers); a combo like DHPP-Lepto covers DHLPP / FVRCP once.",
        "Dates as yyyy-mm-dd; leave a date empty when the document does not print it (never compute it).",
        "Docs: fix the pet names if needed, then set Verified = Y. Only verified docs become gold labels.",
        "Orange = the extractor itself flagged it for review.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 120
    meta = {"targets": targets, "folder": str(Path(a.folder).resolve()), "seed": a.seed, "files": [f.name for f in files]}
    (out / "sample.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    wb.save(out / "review.xlsx")
    print(f"wrote {out / 'review.xlsx'} and {len(files)} predictions in {preds}")
    return 0


def build(a) -> int:
    out = Path(a.out)
    targets = json.loads((out / "sample.json").read_text(encoding="utf-8"))["targets"]
    wb = openpyxl.load_workbook(out / "review.xlsx", read_only=True)
    rows = lambda ws: [dict(zip(EV_COLS if ws.title == "Events" else DOC_COLS, r))  # noqa: E731
                       for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]
    ev_by_doc = defaultdict(list)
    for r in rows(wb["Events"]):
        ev_by_doc[str(r["Doc"]).strip()].append(r)
    gold = Path(a.gold)
    gold.mkdir(exist_ok=True)
    n, unchecked = 0, []

    def iso(v):
        return v.date().isoformat() if hasattr(v, "date") else (str(v).strip() or None) if v is not None else None

    for d in rows(wb["Docs"]):
        if str(d["Verified (Y)"] or "").strip().upper() != "Y":
            continue
        doc = str(d["Doc"]).strip()
        names = [s.strip() for s in str(d["Pets (as printed, ';'-separated)"] or "").split(";") if s.strip()]
        pets = {nm: [] for nm in names}
        for r in ev_by_doc[doc]:
            status = str(r["Status"] or "").strip()
            if not status:
                unchecked.append(doc)
            if status.lower() == "delete":
                continue
            pets.setdefault(str(r["Pet"] or "").strip(), []).append(
                {"target": str(r["Requested name"]).strip(), "dateAdministered": iso(r["Date given (yyyy-mm-dd)"]),
                 "dueDate": iso(r["Due date (yyyy-mm-dd)"])})
        body = {"targets": targets, "source": d["File"],
                "pets": [{"name": nm, "events": es} for nm, es in pets.items()]}
        (gold / f"{doc}.json").write_text(json.dumps(body, indent=1), encoding="utf-8")
        n += 1
    print(f"wrote {n} gold label(s) to {gold}")
    if unchecked:
        print(f"note: verified docs with rows that have no Status: {sorted(set(unchecked))}", file=sys.stderr)
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("draft", help="sample documents, extract them, write the review workbook")
    d.add_argument("folder")
    d.add_argument("--n", type=int, default=50)
    d.add_argument("--seed", type=int, default=0)
    d.add_argument("--targets", default=DEFAULT_TARGETS)
    d.add_argument("--out", default="gold_draft")
    b = sub.add_parser("build", help="verified rows of review.xlsx -> gold/<doc>.json")
    b.add_argument("--out", default="gold_draft")
    b.add_argument("--gold", default="gold")
    a = ap.parse_args()
    return draft(a) if a.cmd == "draft" else build(a)


if __name__ == "__main__":
    raise SystemExit(main())
