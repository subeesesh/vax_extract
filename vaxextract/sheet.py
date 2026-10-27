"""Write extraction results into the verification workbook (VAX_SHEET, e.g. vax_verify_failures.xlsx).

One row per (pet x requested vaccine name) in the sheet's existing columns; with no requested names (extraction
only), one row per printed vaccination. Rows already in the sheet for the
same document (Document_URL) are replaced, so re-running a document never duplicates it.
"""
from __future__ import annotations

import json
import logging
from datetime import datetime
from pathlib import Path

from .schema import Result

log = logging.getLogger(__name__)

COLUMNS = ["PetName", "POName", "Immunization_Type_Name", "created_on", "if_deleted", "Expiration_date", "status",
           "Document_URL", "Results - Comment", "is_verified_by_ai", "Verified By", "final_json"]
WIDTHS = {"PetName": 16, "POName": 20, "Immunization_Type_Name": 18, "created_on": 16, "Expiration_date": 13,
          "status": 13, "Document_URL": 40, "Results - Comment": 50, "Verified By": 24, "final_json": 60}


OCR_NAMES = {"rapid": "RapidOCR", "paddle": "PaddleOCR", "tesseract": "Tesseract"}


def verified_by(model: str, ocr_engine: str) -> str:
    return f"vaxextract ({model} + {OCR_NAMES.get(ocr_engine, ocr_engine)})"


def rows_for(result: Result, document: str, verified_by: str, created: datetime) -> list[dict]:
    d = result.data
    owner = d.owner.name if d.owner else None
    base = {"POName": owner, "created_on": created, "Document_URL": document, "Verified By": verified_by,
            "final_json": json.dumps(result.to_output(), ensure_ascii=False)}  # the client payload, same on every row
    if not d.pets or not any(p.vaccinations or p.printedRows for p in d.pets):
        return [{**base, "PetName": d.pets[0].name if d.pets else None, "status": "Not found",
                 "is_verified_by_ai": False,
                 "Results - Comment": "; ".join(["no vaccination rows extracted", *result.meta.reviewReasons])}]
    if not any(p.vaccinations for p in d.pets):  # nothing requested: no mapping, list the printed rows
        return _printed_rows_for(result, base)
    rows = []
    for p in d.pets:
        for v in p.vaccinations:
            matched = bool(v.matchedComponents)
            ok = matched and not v.needsHumanReview
            notes = []
            if matched:
                notes += [f"printed as '{v.vaccine_name}'", f"given {v.dateAdministered or 'n/a'}"]
                if v.partialMatch:
                    notes.append("partial match")
                if v.expirationDateSource and v.expirationDateSource != "document":
                    notes.append(f"expiry {v.expirationDateSource}")
            else:
                notes.append("not on document")
            rows.append({**base, "PetName": p.name, "Immunization_Type_Name": v.name,
                         "Expiration_date": datetime.fromisoformat(v.expirationDate) if v.expirationDate else None,
                         "status": "Verified" if ok else ("Needs review" if matched else "Not found"),
                         "is_verified_by_ai": ok, "Results - Comment": "; ".join(notes + v.reviewReasons)})
    return rows


def _printed_rows_for(result: Result, base: dict) -> list[dict]:
    rows = []
    for p in result.to_output()["data"]["pets"]:
        for v in p["vaccinations"]:
            ok = not v["needsHumanReview"]
            rows.append({**base, "PetName": p.get("name"), "Immunization_Type_Name": v["name"],
                         "Expiration_date": datetime.fromisoformat(v["expirationDate"]) if v.get("expirationDate") else None,
                         "status": "Extracted" if ok else "Needs review", "is_verified_by_ai": ok,
                         "Results - Comment": f"given {v.get('dateAdministered') or 'n/a'}"})
    return rows


def update_sheet(path: str | Path, result: Result, document: str, verified_by: str = "vaxextract") -> bool:
    """Replace `document`'s rows in the workbook (created if missing). False if it could not be saved."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    path = Path(path)
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(COLUMNS)
    col = {c.value: c.column for c in ws[1] if c.value}
    for k in COLUMNS:
        if k not in col:
            col[k] = ws.max_column + 1
            ws.cell(row=1, column=col[k], value=k)
    url = col["Document_URL"]
    for r in range(ws.max_row, 1, -1):
        if ws.cell(row=r, column=url).value == document:
            ws.delete_rows(r)
    r = ws.max_row + 1
    while r > 2 and all(c.value is None for c in ws[r - 1]):
        r -= 1

    font, flag, miss = Font(name="Arial", size=10), PatternFill("solid", fgColor="FCE4D6"), PatternFill("solid", fgColor="EDEDED")
    for vals in rows_for(result, document, verified_by, datetime.now().replace(microsecond=0)):
        for k in COLUMNS:
            c = ws.cell(row=r, column=col[k], value=vals.get(k))
            c.font, c.alignment = font, Alignment(vertical="top")
            if vals["status"] == "Needs review":
                c.fill = flag
            elif vals["status"] == "Not found":
                c.fill = miss
        ws.cell(row=r, column=col["created_on"]).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=r, column=col["Expiration_date"]).number_format = "yyyy-mm-dd"
        r += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=ws.max_column).column_letter}{max(r - 1, 1)}"
    for k, w in WIDTHS.items():
        ws.column_dimensions[ws.cell(row=1, column=col[k]).column_letter].width = w
    try:
        wb.save(path)
    except PermissionError:  # usually the workbook is open in Excel
        log.warning("could not write %s (is it open in Excel?)", path)
        return False
    return True
